"""RU: Извлекает текст и изображения презентаций предмета. EN: Extracts subject slide text and images."""
import argparse
import hashlib
import json
from pathlib import Path
from pptx import Presentation
from PIL import Image, ImageDraw
import io


def shapes_in(shapes):
    """RU: Обходит вложенные фигуры. EN: Walks nested slide shapes."""
    for shape in shapes:
        yield shape
        if shape.shape_type == 6:
            yield from shapes_in(shape.shapes)


def extract(source, destination):
    """RU: Сохраняет извлечение с номерами слайдов. EN: Saves extraction with slide numbers."""
    presentation = Presentation(str(source))
    destination.mkdir(parents=True, exist_ok=True)
    pages = []
    visuals = []
    for number, slide in enumerate(presentation.slides, 1):
        texts, images = [], []
        for shape in shapes_in(slide.shapes):
            if shape.has_text_frame:
                texts.append(shape.text)
            if shape.has_table:
                texts.extend(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows)
            if shape.shape_type == 13:
                image_id = f'slide-{number:03d}-image-{len(images)+1:02d}'
                image_file = destination / f'{image_id}.{shape.image.ext}'
                image_file.write_bytes(shape.image.blob)
                images.append(image_file.name)
                try:
                    image = Image.open(io.BytesIO(shape.image.blob)).convert('RGB')
                    visuals.append((image_id, image))
                except Exception:
                    pass
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide and slide.notes_slide.notes_text_frame else ''
        pages.append({'page': number, 'text': '\n'.join(texts), 'notes': notes, 'images': images})
    for start in range(0, len(visuals), 6):
        sheet = Image.new('RGB', (1600, 1800), 'white')
        draw = ImageDraw.Draw(sheet)
        for offset, (label, image) in enumerate(visuals[start:start+6]):
            x, y = (offset % 2)*800, (offset // 2)*600
            image.thumbnail((780, 565))
            sheet.paste(image, (x, y+30))
            draw.text((x+10, y+8), label, fill='black')
        sheet.save(destination / f'contact-{start//6+1:02d}.jpg', quality=90)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    record = {'source': source.name, 'sha256': digest, 'pages': pages}
    (destination / 'extracted.json').write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding='utf-8')
    (destination / 'TEXT.md').write_text('\n\n'.join(f'## Slide {page["page"]}\n\n{page["text"]}\n\nImages: {", ".join(page["images"])}\n\nNotes: {page["notes"]}' for page in pages), encoding='utf-8')
    return {'source': source.name, 'sha256': digest, 'slides': len(pages), 'images': len(visuals), 'extraction': str(destination / 'extracted.json')}


def main():
    """RU: Создаёт инвентарь предмета и свежие извлечения. EN: Builds subject inventory and fresh extracts."""
    parser = argparse.ArgumentParser()
    parser.add_argument('source', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    inventory, known = [], {}
    presentation_index = 0
    for source in sorted(args.source.iterdir()):
        if not source.is_file():
            continue
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        if digest in known:
            inventory.append({'source': source.name, 'sha256': digest, 'duplicateOf': known[digest]})
            continue
        known[digest] = source.name
        if source.suffix.lower() == '.pptx':
            import re
            presentation_index += 1
            match = re.search(r'лекція\s*№?\s*(\d+)', source.name, flags=re.IGNORECASE)
            key = f'lecture-{int(match[1]):02d}' if match else f'presentation-{presentation_index:02d}'
            record = extract(source, args.output / key)
            record['key'] = key
            record['status'] = 'extracted_not_read'
            inventory.append(record)
        else:
            inventory.append({'source': source.name, 'sha256': digest, 'status': 'supplement_not_read'})
    (args.output / 'inventory.json').write_text(json.dumps(inventory, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(inventory, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
